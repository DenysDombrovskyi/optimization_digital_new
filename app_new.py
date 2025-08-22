import streamlit as st
import pandas as pd
import numpy as np
import io
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import numbers
from pulp import *

st.set_page_config(page_title="Digital Split Optimizer", layout="wide")

# ==== Брендинг dentsu X ====
col1, col2 = st.columns([1,5])
with col1:
    st.image("d3d64b71-7d13-456d-8d93-3c3f491f81ec.png", width=140)  # логотип dentsu X
with col2:
    st.markdown(
        "<h1 style='margin-bottom:0;'>Digital Split Optimizer – Гібридний спліт</h1>"
        "<p style='color:gray; font-size:18px;'>powered by dentsu X</p>",
        unsafe_allow_html=True
    )

st.markdown("---")

# ==== Налаштування ====
optimization_goal = st.radio(
    "Оберіть мету оптимізації:",
    ('Максимізація охоплення', 'Мінімізація бюджету'),
    horizontal=True
)
st.markdown("---")

num_instruments = st.number_input("Кількість інструментів:", min_value=1, max_value=50, value=25, step=1)
total_audience = st.number_input("Загальний розмір потенційної аудиторії:", value=50000, step=1000)

if optimization_goal == 'Максимізація охоплення':
    total_budget = st.number_input("Заданий бюджет ($):", value=100000, step=1000)
else:
    reach_target_pct = st.number_input("Бажаний відсоток охоплення (%):", min_value=1, max_value=100, value=50, step=1)
    
# ==== Початкові дані ====
if "df" not in st.session_state or len(st.session_state.df) != num_instruments:
    st.session_state.df = pd.DataFrame({
        "Instrument": [f"Instrument {i+1}" for i in range(num_instruments)],
        "CPM": [50 + i*2 for i in range(num_instruments)],
        "Freq": [1.0 + 0.05*i for i in range(num_instruments)],
        "MinShare": [0.01 for _ in range(num_instruments)],
        "MaxShare": [0.15 for i in range(num_instruments)]
    })

df = st.session_state.df.copy()

# ==== Форма редагування ====
with st.form("instrument_form"):
    st.subheader("Редагування даних інструментів")
    for i in range(len(df)):
        st.markdown(f"**Інструмент {i+1}**")
        col1, col2, col3, col4, col5 = st.columns(5)
        with col1:
            df.loc[i, "Instrument"] = st.text_input(f"Назва", value=df.loc[i, "Instrument"], key=f"name_{i}")
        with col2:
            df.loc[i, "CPM"] = st.number_input(f"CPM", value=float(df.loc[i, "CPM"]), step=1.0, key=f"cpm_{i}")
        with col3:
            df.loc[i, "Freq"] = st.number_input(f"Frequency", value=float(df.loc[i, "Freq"]), step=0.01, key=f"freq_{i}")
        with col4:
            df.loc[i, "MinShare"] = st.number_input(f"Min Share", value=float(df.loc[i, "MinShare"]), step=0.01, min_value=0.0, max_value=1.0, key=f"min_share_{i}")
        with col5:
            df.loc[i, "MaxShare"] = st.number_input(f"Max Share", value=float(df.loc[i, "MaxShare"]), step=0.01, min_value=0.0, max_value=1.0, key=f"max_share_{i}")
        st.markdown("---")
    submitted = st.form_submit_button("Перерахувати спліт")

if submitted:
    st.session_state.df = df.copy()
    df["Efficiency"] = 1000 / (df["CPM"] * df["Freq"])
    df["CPR"] = np.where(df["Freq"] != 0, df["CPM"] / 1000 * df["Freq"], np.inf)
    
    def total_reach(budgets, df_current, total_audience_val):
        budgets_array = np.array(budgets)
        impressions = np.where(df_current["CPM"].values != 0, budgets_array / df_current["CPM"].values * 1000, 0)
        reach_i = np.where(
            df_current["Freq"].values != 0, 
            impressions / df_current["Freq"].values / total_audience_val, 
            0
        )
        reach_i = np.clip(reach_i, 0.0, 1.0)
        return 1 - np.prod(1 - reach_i)

    # ==== Мінімізація бюджету ====
    if optimization_goal == 'Мінімізація бюджету':
        st.subheader(f"Результат: **Мінімізація бюджету**")
        
        if reach_target_pct <= 0:
            st.error("Цільовий відсоток охоплення має бути більше 0%.")
        else:
            reach_target_people = total_audience * (reach_target_pct / 100)
            model = LpProblem("Minimize_Budget_for_Reach_Target", LpMinimize)
            budget_vars = LpVariable.dicts("Budget", df['Instrument'], lowBound=0)
            model += lpSum([budget_vars[instrument] for instrument in df['Instrument']]), "Total Budget"

            unique_reach_per_instrument_terms = []
            for _, row in df.iterrows():
                if row['Freq'] != 0 and row['CPM'] != 0:
                    unique_reach_per_instrument_terms.append(budget_vars[row['Instrument']] / row['CPM'] * 1000 / row['Freq'])
                else:
                    unique_reach_per_instrument_terms.append(0) 

            model += lpSum(unique_reach_per_instrument_terms) >= reach_target_people, "Total Unique Reach Constraint"

            total_budget_lp_var = lpSum([budget_vars[instrument] for instrument in df['Instrument']])
            for _, row in df.iterrows():
                model += budget_vars[row['Instrument']] >= row['MinShare'] * total_budget_lp_var
                model += budget_vars[row['Instrument']] <= row['MaxShare'] * total_budget_lp_var

            status = model.solve()
            
            if status == LpStatusOptimal:
                final_total_budget = value(model.objective)
                df_result = df.copy()
                df_result['Budget'] = [value(budget_vars[i]) for i in df_result['Instrument']]
                df_result["BudgetSharePct"] = df_result["Budget"] / final_total_budget if final_total_budget > 0 else 0
                df_result["Impressions"] = np.where(df_result["CPM"] != 0, df_result["Budget"] / df_result["CPM"] * 1000, 0)
                df_result["Unique Reach (People)"] = np.where(df_result["Freq"] != 0, df_result["Impressions"] / df_result["Freq"], 0)
                df_result["ReachPct"] = df_result["Unique Reach (People)"] / total_audience
                df_result["ReachPct"] = df_result["ReachPct"].clip(upper=1.0) 
                total_reach_prob = total_reach(df_result["Budget"].values, df_result, total_audience)
                
                st.success(f"Оптимальне рішення знайдено! 🎉")
                st.write(f"Мінімальний **Бюджет** для охоплення **{total_reach_prob*100:.2f}%** аудиторії: **{final_total_budget:,.2f} $**")
                st.dataframe(df_result)
            else:
                st.warning(f"Не знайдено оптимального рішення. Статус: **{LpStatus[status]}**")

    # ==== Максимізація охоплення ====
    elif optimization_goal == 'Максимізація охоплення':
        st.subheader("Результат: **Максимізація охоплення** (Лінійне програмування)")
        model = LpProblem("Maximize_Reach_LP_Approximation", LpMaximize)
        budget_vars = LpVariable.dicts("Budget", df['Instrument'], lowBound=0)
        objective_terms = []
        for _, row in df.iterrows():
            if row['Freq'] != 0 and row['CPM'] != 0:
                objective_terms.append(budget_vars[row['Instrument']] / (row['CPM'] * row['Freq']) * 1000)
            else:
                objective_terms.append(0) 
        model += lpSum(objective_terms), "Total Unique Reach"
        model += lpSum([budget_vars[instrument] for instrument in df['Instrument']]) <= total_budget
        for _, row in df.iterrows():
            model += budget_vars[row['Instrument']] >= row['MinShare'] * total_budget
            model += budget_vars[row['Instrument']] <= row['MaxShare'] * total_budget
        status = model.solve()
        
        if status == LpStatusOptimal:
            maximized_sum_of_unique_reach_people = value(model.objective) 
            df_result = df.copy()
            df_result['Budget'] = [value(budget_vars[i]) for i in df_result['Instrument']]
            total_budget_actual_spent = df_result['Budget'].sum()
            df_result["BudgetSharePct"] = df_result["Budget"] / total_budget_actual_spent if total_budget_actual_spent > 0 else 0
            df_result["Impressions"] = np.where(df_result["CPM"] != 0, df_result["Budget"] / df_result["CPM"] * 1000, 0)
            df_result["Unique Reach (People)"] = np.where(df_result["Freq"] != 0, df_result["Impressions"] / df_result["Freq"], 0)
            df_result["ReachPct"] = df_result["Unique Reach (People)"] / total_audience
            df_result["ReachPct"] = df_result["ReachPct"].clip(upper=1.0)
            total_reach_prob = total_reach(df_result["Budget"].values, df_result, total_audience)
            
            st.success(f"Оптимальне рішення знайдено! 🎉")
            st.write(f"Максимізоване **Охоплення** (лінійно): **{maximized_sum_of_unique_reach_people:,.0f} людей**")
            st.write(f"Фактичне **Охоплення**: **{total_reach_prob*100:.2f}%** аудиторії")
            st.write(f"Витрачений **Бюджет**: **{total_budget_actual_spent:,.2f} $** (з доступних {total_budget:,.2f} $)")
            st.dataframe(df_result)
        else:
            st.warning(f"Не знайдено оптимального рішення. Статус: **{LpStatus[status]}**")

    # ==== Завантаження результатів у Excel ====
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Гібридний спліт"
    
    excel_cols = ["Instrument", "CPM", "CPR", "Freq", "MinShare", "MaxShare", "Budget", "BudgetSharePct", "Impressions", "Unique Reach (People)", "ReachPct"]

    df_to_save = df_result[excel_cols].copy()
    final_total_budget = df_result["Budget"].sum()
    final_total_reach_prob = total_reach(df_result["Budget"].values, df_result, total_audience)
    final_total_reach_people_linear_sum = df_result["Unique Reach (People)"].sum()
    
    df_to_save.loc[len(df_to_save)] = [
        "TOTAL", np.nan, np.nan, np.nan, 
        df_result['MinShare'].sum(),
        df_result['MaxShare'].sum(),
        final_total_budget, 
        1.0,
        df_result["Impressions"].sum(), 
        final_total_reach_people_linear_sum, 
        f"{final_total_reach_prob*100:.2f}%"
    ]

    for r in dataframe_to_rows(df_to_save, index=False, header=True):
        ws.append(r)
    
    budget_share_pct_col = excel_cols.index("BudgetSharePct") + 1
    for row in ws.iter_rows(min_row=2, max_row=1+len(df_result), min_col=budget_share_pct_col, max_col=budget_share_pct_col):
        for cell in row:
            cell.number_format = '0.00%'
            
    chart = BarChart()
    chart.type = "col"
    chart.title = "Бюджет по інструментам (%)"
    chart.y_axis.title = "Budget Share (%)"
    chart.x_axis.title = "Інструменти"
    data = Reference(ws, min_col=budget_share_pct_col, min_row=1, max_row=1+len(df_result))
    categories = Reference(ws, min_col=1, min_row=2, max_row=1+len(df_result))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, f"L{budget_share_pct_col}") 

    wb.save(output)
    output.seek(0)
    
    st.download_button(
        label="Завантажити результати (Excel)",
        data=output,
        file_name="Digital_Split_Hybrid.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
